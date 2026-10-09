"""Таблица услуг для заказчика: категории сайта → разделы → группы → услуги, наши тексты и цены.

Собирается из тех же данных, что сайт:
  content.json        — шесть категорий и какие листы прайса в них входят (наши названия разделов)
  prices.json         — услуги, время, цены по классам (импорт из xlsx заказчика)
  service-texts.json  — наша редактура: название, «что это», «что входит», условия

    python scripts/export-services-table.py          → ../Услуги_и_цены.xlsx

Прайс обновился — запустить заново. Колонка «Код» — ключ услуги (лист:услуга), по нему
правки заказчика переносятся обратно в service-texts.json; её не менять."""
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'src' / 'data'
OUT = ROOT.parent / 'Услуги_и_цены.xlsx'

prices = json.loads((DATA / 'prices.json').read_text(encoding='utf-8'))
texts = json.loads((DATA / 'service-texts.json').read_text(encoding='utf-8'))
content = json.loads((DATA / 'content.json').read_text(encoding='utf-8'))

# lib/typo.js → sentence(): капс прайса в обычный регистр, бренды и аббревиатуры как есть
KEEP = {'KRYTEX': 'Krytex', 'SUNTEK': 'SunTek', 'WEMATEC': 'Wematec', 'SPECTROLL': 'Spectroll',
        'SIO2': 'SiO2', 'NANO': 'Nano', 'TOPCOAT': 'TopCoat'}


def sentence(text):
    if not text:
        return text or ''
    out = []
    for word in re.split(r'(\s+)', str(text).strip()):
        core = re.sub(r'[^\wА-Яа-яЁё]', '', word)
        up = core.upper()
        if core and KEEP.get(up) and core == up:
            word = word.replace(core, KEEP[up])
        elif core and len(core) > 3 and core == up:
            word = word.lower()
        out.append(word)
    s = ''.join(out)
    return s[:1].upper() + s[1:]


def service_text(sheet_id, s):
    own = texts.get(sheet_id, {}).get(s['id'], {})
    return {
        'name': own.get('name') or sentence(s['name']),
        'lead': own.get('lead') or ' '.join(s.get('desc') or []),
        'includes': own.get('includes') or s.get('steps') or [],
        'terms': own.get('terms') or s.get('notes') or [],
        'own': bool(own),
    }


sheets = {sh['sheet']: sh for sh in prices['categories']}
foot = {f['marker']: f['text'] for f in prices.get('footnotes', [])}

HEAD = ['Категория', 'Раздел', 'Группа', 'Услуга на сайте', 'Как в прайсе', 'Что это (наш текст)',
        'Что входит', 'Условия', 'Время', '1-й класс', '2-й класс', '3-й класс', '4-й класс', '5-й класс',
        'Примечание к цене', 'Комментарий', 'Код (не менять)']
WIDTH = [16, 22, 22, 30, 30, 46, 52, 36, 11, 11, 11, 11, 11, 11, 30, 30, 26]
PRICE_COL = HEAD.index('1-й класс') + 1

wb = openpyxl.Workbook()
info = wb.active
info.title = 'Как читать'
rows = [
    ['Услуги и цены сайта'],
    [],
    ['Категория — плитка на сайте (шесть штук), Раздел — лист вашего прайса под нашим названием,'],
    ['Группа — подзаголовок внутри раздела, если он есть в прайсе.'],
    ['«Услуга на сайте» и «Что это / Что входит / Условия» — наши тексты; «Как в прайсе» — ваше название для сверки.'],
    ['Цены — из вашего прайса, по классу машины. «от» — цена от, уточняется после осмотра.'],
    ['Правки и замечания пишите в жёлтую колонку «Комментарий». Колонку «Код» не меняйте — по ней мы переносим правки на сайт.'],
]
for r in rows:
    info.append(r)
info['A1'].font = Font(bold=True, size=14)
info.column_dimensions['A'].width = 110

ws = wb.create_sheet('Услуги')
ws.append(HEAD)
bold = Font(bold=True)
head_fill = PatternFill('solid', fgColor='DDDDDD')
yellow = PatternFill('solid', fgColor='FFF2CC')
cat_fill = PatternFill('solid', fgColor='EFEFEF')
wrap = Alignment(wrap_text=True, vertical='top')
for i, c in enumerate(ws[1], 1):
    c.font = bold
    c.fill = yellow if HEAD[i - 1] == 'Комментарий' else head_fill
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = WIDTH[i - 1]

total = 0
for cat in content['services']['categories']:
    # строка-заголовок категории — чтобы таблица читалась блоками
    ws.append([cat['title']])
    for c in ws[ws.max_row]:
        c.fill = cat_fill
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    for ref in cat['sheets']:
        sh = sheets[ref['sheet']]
        labels = sh['classLabels']
        odd = len(labels) != 5  # мототехника: свои два класса
        for g in sh['groups']:
            for s in g['services']:
                t = service_text(sh['id'], s)
                notes = []
                if odd:
                    notes.append('Классы мото: ' + ', '.join(f'{i + 1} — {l}' for i, l in enumerate(labels)))
                if s.get('footnote'):
                    notes.append(foot.get(s['footnote'], s['footnote']))
                row = [cat['title'], ref.get('title') or sh['title'], sentence(g.get('title') or ''), t['name'],
                       sentence(s['name']), t['lead'], '\n'.join(f'• {x}' for x in t['includes']),
                       ' '.join(t['terms']), (s.get('duration') or '').strip()]
                cells = []
                for label in labels[:5]:
                    p = s['prices'].get(label)
                    v = (p.get('value') if p.get('value') is not None else p.get('min')) if p else None
                    cells.append((v, p.get('kind') if p else None))
                cells += [(None, None)] * (5 - len(cells))
                row += [v if v is not None else 'по запросу' for v, _ in cells]
                row += ['\n'.join(notes), None, f"{sh['id']}:{s['id']}"]
                ws.append(row)
                r = ws.max_row
                for c in ws[r]:
                    c.alignment = wrap
                ws.cell(r, 4).font = bold
                ws.cell(r, HEAD.index('Комментарий') + 1).fill = yellow
                for k, (v, kind) in enumerate(cells):
                    if v is not None:
                        ws.cell(r, PRICE_COL + k).number_format = ('"от "' if kind == 'from' else '') + '# ##0 ₽'
                total += 1

ws.freeze_panes = 'E2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(HEAD))}{ws.max_row}'
wb.save(OUT)
print(f'{OUT.name}: {total} услуг в {len(content["services"]["categories"])} категориях')
