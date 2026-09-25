"""Сводная таблица машин для заказчика: Машины_и_классы.xlsx + .html в корне проекта.

Запуск из akb-bu: npm run cars-table. Данные берёт scripts/export-cars-table.mjs
из того же справочника, что поиск на сайте (prices.json + car-aliases + car-extra).
Жёлтые строки — машин нет в прайсе, класс предложен по похожей модели.
"""
import os
import sys

os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')  # openpyxl тянет numpy, а тот без этого падает на Windows
sys.stdout.reconfigure(encoding='utf-8')

import html
import json
import subprocess
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = HERE.parent.parent
TODAY = date.today().strftime('%d.%m.%Y')

raw = subprocess.run(['node', str(HERE / 'export-cars-table.mjs')], capture_output=True, check=True, cwd=HERE.parent).stdout
rows = json.loads(raw.decode('utf-8'))
rows.sort(key=lambda r: (r['brand'].lower(), r['cls'], r['model'].lower()))

TITLE = 'Машины на сайте dssever.ru и их классы'
HOW = [
    'Белые строки — машины из вашего прайса (лист «Классификация машин»), класс ваш.',
    'Жёлтые строки — машин нет в прайсе. Класс мы предложили по похожей модели из прайса (колонка «Похожа на»).',
    'Проверьте жёлтые строки: если класс другой — впишите верный в колонку «Ваш класс», вопросы — в «Комментарий».',
    'Лучше всего добавить эти машины в лист «Классификация машин» — тогда класс будет ваш, а не по аналогии.',
]
HEAD = ['№', 'Марка', 'Модель', 'Класс', 'Откуда класс', 'Похожа на (из прайса)', 'Как ещё ищут на сайте', 'Ваш класс', 'Комментарий']

# ---------- xlsx ----------
wb = Workbook()
ws = wb.active
ws.title = 'Машины'
ws['A1'] = TITLE
ws['A1'].font = Font(size=14, bold=True)
for i, line in enumerate(HOW, start=2):
    ws.cell(row=i, column=1, value=line)
ws.cell(row=len(HOW) + 2, column=1, value=f'Выгрузка {TODAY}: моделей {len(rows)}, из них по аналогии {sum(1 for r in rows if r["like"])}.')
top = len(HOW) + 4
ws.append([])
for c, name in enumerate(HEAD, start=1):
    cell = ws.cell(row=top, column=c, value=name)
    cell.font = Font(bold=True)
    cell.fill = PatternFill('solid', fgColor='E7E6E3')
    cell.alignment = Alignment(vertical='center', wrap_text=True)

yellow = PatternFill('solid', fgColor='FFF2B3')
mine = PatternFill('solid', fgColor='DDEBF7')
for n, r in enumerate(rows, start=1):
    line = top + n
    values = [n, r['brand'], r['model'], r['cls'], 'по аналогии' if r['like'] else 'прайс', r['like'], ', '.join(r['alt']), None, None]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=line, column=c, value=v)
        cell.alignment = Alignment(vertical='top', wrap_text=c == 7)
        if r['like'] and c <= 7:
            cell.fill = yellow
    ws.cell(row=line, column=8).fill = mine
    ws.cell(row=line, column=9).fill = mine

for col, width in zip('ABCDEFGHI', [6, 16, 26, 8, 14, 30, 48, 11, 36]):
    ws.column_dimensions[col].width = width
ws.freeze_panes = ws.cell(row=top + 1, column=3)
ws.auto_filter.ref = f'A{top}:I{top + len(rows)}'
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True,
                    error='Класс — число от 1 до 5', errorTitle='Класс')
ws.add_data_validation(dv)
dv.add(f'H{top + 1}:H{top + len(rows)}')

# сводка по классам
sv = wb.create_sheet('Сводка')
sv.append(['Класс', 'Всего моделей', 'Из прайса', 'По аналогии'])
for c in range(1, 5):
    sv.cell(row=1, column=c).font = Font(bold=True)
for k in range(1, 6):
    part = [r for r in rows if r['cls'] == k]
    sv.append([f'{k}-й', len(part), sum(1 for r in part if not r['like']), sum(1 for r in part if r['like'])])
sv.append(['Итого', len(rows), sum(1 for r in rows if not r['like']), sum(1 for r in rows if r['like'])])
sv.cell(row=7, column=1).font = Font(bold=True)
sv.append([])
sv.append(['Марки, которых нет в прайсе целиком'])
sv.cell(row=9, column=1).font = Font(bold=True)
brands = Counter(r['brand'] for r in rows)
price_brands = {r['brand'] for r in rows if not r['like']}
for b in sorted(brands, key=str.lower):
    if b not in price_brands:
        sv.append([b, brands[b]])
for col, width in zip('ABCD', [34, 16, 12, 14]):
    sv.column_dimensions[col].width = width

xlsx = OUT / 'Машины_и_классы.xlsx'
wb.save(xlsx)

# ---------- html ----------
e = html.escape
body = '\n'.join(
    f'<tr class="{"like" if r["like"] else ""}" data-cls="{r["cls"]}">'
    f'<td>{e(r["brand"])}</td><td>{e(r["model"])}</td><td class="num">{r["cls"]}</td>'
    f'<td>{"по аналогии" if r["like"] else "прайс"}</td><td>{e(r["like"])}</td>'
    f'<td class="alt">{e(", ".join(r["alt"]))}</td></tr>'
    for r in rows
)
page = f'''<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Машины и классы</title>
<style>
  :root {{ --bg: #fff; --fg: #141414; --line: #dcdad6; --head: #efeeeb; --like: #fff2b3; --field: #fff; }}
  @media (prefers-color-scheme: dark) {{ :root {{ --bg: #141414; --fg: #fff; --line: #3a3a3a; --head: #222; --like: #4a3f10; --field: #1c1c1c; }} }}
  body {{ margin: 0; background: var(--bg); color: var(--fg); font: 15px/1.45 system-ui, sans-serif; }}
  .wrap {{ max-width: 1240px; margin: 0 auto; padding: 24px 16px 48px; }}
  h1 {{ font-size: 26px; margin: 0 0 12px; max-width: 30ch; }}
  .how {{ max-width: 76ch; margin: 0 0 20px; padding: 0 0 0 18px; }}
  .how li {{ margin: 0 0 6px; }}
  mark {{ background: var(--like); color: inherit; padding: 0 3px; }}
  .bar {{ display: flex; flex-wrap: wrap; gap: 12px 20px; align-items: center; margin: 0 0 12px; }}
  input[type=search], select {{ font: inherit; color: var(--fg); background: var(--field); border: 1px solid var(--line); border-radius: 8px; padding: 8px 12px; }}
  input[type=search] {{ flex: 1 1 260px; }}
  .scroll {{ overflow-x: auto; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th, td {{ text-align: left; vertical-align: top; padding: 7px 10px; border-bottom: 1px solid var(--line); }}
  th {{ position: sticky; top: 0; background: var(--head); font-weight: 600; white-space: nowrap; }}
  tr.like td {{ background: var(--like); }}
  td.num {{ font-variant-numeric: tabular-nums; }}
  td.alt {{ min-width: 220px; }}
</style></head><body><div class="wrap">
<h1>{TITLE}</h1>
<ul class="how">{"".join(f"<li>{e(x)}</li>" for x in HOW[:2])}
<li>Править удобнее в файле <b>Машины_и_классы.xlsx</b> — там есть колонки «Ваш класс» и «Комментарий».</li></ul>
<div class="bar">
  <input type="search" id="q" placeholder="Марка или модель" autocomplete="off">
  <select id="src"><option value="">Все машины</option><option value="like">Только по аналогии</option><option value="price">Только из прайса</option></select>
  <select id="cls"><option value="">Все классы</option>{"".join(f'<option value="{k}">{k}-й класс</option>' for k in range(1, 6))}</select>
  <span id="count"></span>
</div>
<div class="scroll"><table>
<thead><tr><th>Марка</th><th>Модель</th><th>Класс</th><th>Откуда класс</th><th>Похожа на (из прайса)</th><th>Как ещё ищут на сайте</th></tr></thead>
<tbody id="rows">{body}</tbody></table></div>
<p>Выгрузка {TODAY}.</p>
</div>
<script>
  const rows = [...document.querySelectorAll('#rows tr')];
  const q = document.getElementById('q'), src = document.getElementById('src'), cls = document.getElementById('cls');
  const count = document.getElementById('count');
  function apply() {{
    const t = q.value.trim().toLowerCase();
    let n = 0;
    for (const r of rows) {{
      const ok = (!t || r.textContent.toLowerCase().includes(t))
        && (!src.value || (src.value === 'like') === r.classList.contains('like'))
        && (!cls.value || r.dataset.cls === cls.value);
      r.hidden = !ok; if (ok) n++;
    }}
    count.textContent = 'Показано: ' + n + ' из ' + rows.length;
  }}
  [q, src, cls].forEach((el) => el.addEventListener('input', apply));
  apply();
</script></body></html>'''
(OUT / 'Машины_и_классы.html').write_text(page, encoding='utf-8')
print(f'✓ {xlsx.name} и Машины_и_классы.html: моделей {len(rows)}, по аналогии {sum(1 for r in rows if r["like"])}')
